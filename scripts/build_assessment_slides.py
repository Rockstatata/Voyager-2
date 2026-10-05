"""Thirteen-slide assessment deck with one layout shared by PDF and PPTX.

Text remains editable in PowerPoint. Equations and engineering diagrams are
high-resolution figures, screenshots have source provenance, and every technical
slide has derivations, limitations and code pointers in its speaker notes.
"""
import io
import json
import re
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
from matplotlib import mathtext, font_manager
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Pt
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from update_showcase_scope_slides import CARDS

ROOT=Path(__file__).resolve().parents[1]
REPORT=ROOT/'docs/report'
FIG=REPORT/'figures'
TECH=FIG/'technical'
notes_source=(REPORT/'Presentation_Notes_Source.md').read_text(encoding='utf-8')
NOTES={int(m.group(1)):m.group(2).strip() for m in re.finditer(r'^# Slide (\d+):[^\n]+\n(.*?)(?=^# Slide \d+:|\Z)',notes_source,re.M|re.S)}
if len(NOTES)!=13:raise ValueError('Expected 13 source note sections')
W,H=960,540
INK='14283B'; MUTED='546879'; TEAL='167F82'; BLUE='0F4D92'; PALE='EDF3F6'; WHITE='FFFFFF'

for name,file in [('Arial','arial.ttf'),('Arial-Bold','arialbd.ttf')]:
    pdfmetrics.registerFont(TTFont(name,'C:/Windows/Fonts/'+file))

def wrapped(text,width,size,bold=False):
    result=[]; face='Arial-Bold' if bold else 'Arial'
    for para in text.split('\n'):
        line=''
        for word in para.split():
            candidate=(line+' '+word).strip()
            if pdfmetrics.stringWidth(candidate,face,size)>width and line:
                result.append(line);line=word
            else:line=candidate
        result.append(line)
    return result

class Deck:
    def __init__(self):
        self.ppt=Presentation();self.ppt.slide_width=Pt(W);self.ppt.slide_height=Pt(H)
        self.pdf=canvas.Canvas(str(REPORT/'Voyager_2_Presentation.pdf'),pagesize=(W,H))
        self.pdf.setTitle('Voyager 2 | Graphics laboratory presentation | 13 slides')
        self.pdf.setAuthor('Sarwad Hasan Siddiqui')
        self.notes=[];self.layout=[]

    def rect(self,x,y,w,h,color=PALE):
        s=self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h))
        s.fill.solid();s.fill.fore_color.rgb=RGBColor.from_string(color);s.line.fill.background()
        self.pdf.setFillColor('#'+color);self.pdf.rect(x,H-y-h,w,h,fill=1,stroke=0)

    def text(self,text,x,y,w,size=18,bold=False,color=INK,leading=None):
        leading=leading or size*1.25
        lines=wrapped(text,w,size,bold);height=len(lines)*leading+4
        if y+height>535:raise ValueError(f'Slide {len(self.notes)} text overflows: {text[:60]}')
        self.layout.append({'slide':len(self.notes),'kind':'text','text':text,'box':[x,y,w,height],'font_pt':size})
        box=self.slide.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(height));tf=box.text_frame
        tf.margin_top=tf.margin_bottom=tf.margin_left=tf.margin_right=0;tf.word_wrap=False
        for i,line in enumerate(lines):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line
            p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
            p.space_after=Pt(0);p.line_spacing=Pt(leading)
        self.pdf.setFillColor('#'+color);self.pdf.setFont('Arial-Bold' if bold else 'Arial',size)
        for i,line in enumerate(lines):self.pdf.drawString(x,H-y-size-i*leading,line)
        return height

    def pic(self,path,x,y,w,h,crop=None):
        path=Path(path)
        im=Image.open(path).convert('RGB')
        if crop:im=im.crop(crop)
        scale=min(w/im.width,h/im.height);pw,ph=im.width*scale,im.height*scale
        px=x+(w-pw)/2;py=y+(h-ph)/2
        self.layout.append({'slide':len(self.notes),'kind':'image','source':str(path.relative_to(ROOT)),
                            'crop':crop,'box':[px,py,pw,ph]})
        buf=io.BytesIO();im.save(buf,format='PNG');buf.seek(0)
        self.slide.shapes.add_picture(buf,Pt(px),Pt(py),width=Pt(pw),height=Pt(ph))
        self.pdf.drawImage(ImageReader(im),px,H-py-ph,pw,ph)

    def eq(self,formula,x,y,w,size=19):
        formula=re.sub(r'\\mathbf\s+([A-Za-z])',r'\\mathbf{\1}',formula)
        buf=io.BytesIO()
        mathtext.math_to_image('$'+formula+'$',buf,dpi=300,format='png',
                              prop=font_manager.FontProperties(size=size,math_fontfamily='cm'))
        buf.seek(0);im=Image.open(buf);scale=min(72/300,w/im.width)
        pw,ph=im.width*scale,im.height*scale
        self.layout.append({'slide':len(self.notes),'kind':'equation','formula':formula,'box':[x,y,pw,ph]})
        if y+ph>515:raise ValueError('equation outside slide')
        self.slide.shapes.add_picture(buf,Pt(x),Pt(y),width=Pt(pw),height=Pt(ph))
        self.pdf.drawImage(ImageReader(im),x,H-y-ph,pw,ph,mask='auto')
        return ph

    def start(self,title,subtitle,notes):
        notes=NOTES[len(self.notes)+1]
        self.slide=self.ppt.slides.add_slide(self.ppt.slide_layouts[6]);self.notes.append((title,notes))
        self.slide.background.fill.solid();self.slide.background.fill.fore_color.rgb=RGBColor.from_string(WHITE)
        self.rect(0,0,W,6,TEAL)
        self.text(title,30,20,900,29,True)
        self.text(subtitle,32,58,896,14,color=MUTED)
        self.rect(30,503,900,1,'C8D5DD')
        self.text('Sarwad Hasan Siddiqui · 2107006 · CSE 4102 · KUET',32,511,800,10,color=MUTED)
        self.text(f'{len(self.notes):02d} / 13',876,511,65,10,color=MUTED)

    def end(self):
        self.slide.notes_slide.notes_text_frame.text=self.notes[-1][1]
        self.pdf.showPage()

    def save(self):
        assert len(self.notes)==13
        for i,a in enumerate(self.layout):
            for b in self.layout[i+1:]:
                if a['slide']!=b['slide']:continue
                ax,ay,aw,ah=a['box'];bx,by,bw,bh=b['box']
                overlap_x=min(ax+aw,bx+bw)-max(ax,bx)
                overlap_y=min(ay+ah,by+bh)-max(ay,by)
                if overlap_x>3 and overlap_y>3:
                    raise ValueError(f"Slide {a['slide']} content overlaps: {a.get('text',a['kind'])[:55]} / {b.get('text',b['kind'])[:55]}")
        self.pdf.save();self.ppt.save(REPORT/'Voyager_2_Presentation.pptx')
        (REPORT/'Presentation_Notes.md').write_text('\n\n'.join(f'# Slide {i}: {t}\n\n{n}' for i,(t,n) in enumerate(self.notes,1))+'\n',encoding='utf-8')
        (REPORT/'validation/slide-layout.json').write_text(json.dumps(self.layout,indent=2)+'\n',encoding='utf-8')

def main():
    d=Deck()
    d.start('Voyager 2 Solar-System Explorer','01 INTRODUCTION · What the project is and why it was built',
        "")
    d.text('Explore a 3D solar system.\nExplain how it is rendered.',32,101,425,27,True)
    d.text('Replay flybys and inspect the spacecraft.\nFly manually. Compare rendering techniques.',32,177,435,18)
    d.rect(32,244,435,118)
    d.text('26 textured bodies · 4 ring systems\n76 assemblies · 18 hardware targets\n8,500 instanced rocks · dynamic comet',46,258,407,18,True,leading=28)
    d.pic(ROOT/'captures/assessment/technical/voyager_clean.bmp',492,95,436,268)
    d.pic(FIG/'overview.jpg',32,374,282,104)
    d.pic(FIG/'saturn_encounter.jpg',325,374,282,104)
    d.text('Sarwad Hasan Siddiqui | Roll: 2107006\nCSE 4102 | KUET | October 2026',631,382,296,17,True)
    d.end()

    d.start('Explore, control, compare','02 INTRODUCTION · A complete project walkthrough, followed by your recorded video',
        "")
    labels=[('MISSION','Dated launch, four flybys and heliopause.','jupiter_encounter.jpg'),
            ('INSPECTION','Body focus, 18 hardware targets, free camera.','dish.jpg'),
            ('COMPARISON','Three lights, five shading modes, ray tracing.','saturn_traced.jpg')]
    for i,(a,b,c) in enumerate(labels):
        x=32+i*306;d.pic(FIG/c,x,99,286,161);d.text(a,x,272,286,20,True,color=TEAL);d.text(b,x,304,286,17)
    for i,(heading,body) in enumerate(CARDS):
        x=32+i*306
        d.rect(x,373,286,117,PALE)
        d.text(heading,x+12,383,262,15,True,color=TEAL)
        d.text(body,x+12,410,262,14,leading=18)
    d.rect(642,378,286,97,INK)
    d.text('PLAY PROJECT VIDEO',659,392,254,19,True,color=WHITE)
    d.text('Attach your final recording here.',659,428,254,14,color=WHITE)
    d.end()

    d.start('Vertices → indices → triangles → worlds','03 GEOMETRY · Shared sphere construction creates all 26 textured bodies',
        "")
    d.pic(ROOT/'captures/assessment/technical/earth_wire.bmp',32,93,265,180,crop=(160,120,1260,790))
    d.pic(TECH/'index_cell.png',305,95,275,181,crop=(25,90,610,555))
    d.pic(TECH/'body_catalog.png',608,93,320,289)
    d.text('Vertex = position + normal + UV',32,286,550,20,True)
    d.eq(r'\mathbf p=(\sin\phi\cos\theta,\cos\phi,\sin\phi\sin\theta)',32,319,558,19)
    d.eq(r'\phi=\pi i/L,\quad\theta=2\pi j/N,\quad uv=(1-j/N,1-i/L)',32,360,558,18)
    d.text('Cell triangles: (TL, TR, BL) and (TR, BR, BL).\nPoles omit degenerate halves; the UV seam duplicates positions.',32,403,550,16)
    d.rect(608,399,320,85)
    d.text('32 × 64 shared sphere\n2,145 vertices · 3,968 triangles\n11,904 indices · 26 materials/maps',620,410,295,17,True)
    d.end()

    d.start('Voyager: authored from reusable primitives','04 OBJECT CONSTRUCTION · Boxes, cylinders, frustums, rods and a curved parabolic dish',
        "")
    d.pic(TECH/'primitive_profiles.png',32,88,896,178)
    d.text('BOX: 24 vertices / 12 triangles',32,273,292,17,True)
    d.text('CAPS + SIDES: 4n+6 / 4n',338,273,292,17,True)
    d.text('DISH: 590 / 1,152',644,273,284,17,True)
    d.eq(r'y=-D+D(r/R)^2',32,319,286,23)
    d.eq(r'\mathbf n\propto(-2Dx/R^2,1,-2Dz/R^2)',338,319,590,21)
    d.pic(ROOT/'captures/assessment/technical/dish_clean.bmp',32,370,282,119,crop=(300,30,1230,875))
    d.pic(ROOT/'captures/assessment/technical/dish_wire.bmp',325,370,282,119,crop=(300,30,1230,875))
    d.text('76 assemblies · 5,828 triangles\n18 inspectable hardware components\nRods: rotate +Y onto B−A; rebase indices.',631,380,296,17,True)
    d.end()

    d.start('Rings, belts, stars, comet and scene guides','05 OBJECT FAMILIES · Every rendered family has a primitive and a construction rule',
        "")
    d.pic(TECH/'environment.png',32,94,896,151)
    cols=[('RINGS','p = (r cosθ, 0, r sinθ)\nInner/outer edges; reversed back face.\n260 vertices / 256 triangles per band.'),
          ('ROCK FIELDS','28 vertices / 24 triangles per rock.\n8,500 matrices → 3 instanced draws.\nDeterministic placement; static fields.'),
          ('STARS / GUIDES','7,380 points; 96-vertex line circles.\nNo surface triangles for these guides.\nHUD text uses screen-space quads.')]
    for i,(a,b) in enumerate(cols):
        x=32+i*306;d.text(a,x,265,286,19,True,color=TEAL);d.text(b,x,300,286,16)
    d.rect(32,402,896,38)
    d.text('DYNAMIC COMET = sphere nucleus + glow coma + 16-segment cone tail',46,410,868,18,True)
    d.eq(r'\mathbf d_{tail}=\operatorname{normalize}(\mathbf p_{comet}-\mathbf p_{Sun})',46,445,555,20)
    d.text('52 cone vertices / 32 triangles',635,447,275,16)
    d.end()

    d.start('Transforms, hierarchy, scale and cameras','06 TRANSFORMATIONS · Position, pose, size and view are separate operations',
        "")
    d.eq(r'\mathbf p_{clip}=PV M_{world}\mathbf p,\quad M_{world}=M_{parent}TRS',32,101,896,24)
    d.pic(TECH/'transform_hierarchy.png',32,160,560,233)
    d.text('Four camera cases',618,165,308,21,True,color=TEAL)
    d.text('FreeFly: independent motion\nFocus: body-centred orbit\nChase: spacecraft frame\nInspect: hardware close-up',618,205,308,18,leading=31)
    d.eq(r'\mathbf n^{\prime}=\operatorname{normalize}((M_{3\times3}^{-1})^T\mathbf n)',32,410,555,21)
    d.text('Subtract the eye in double precision.\nCompressed scale keeps bodies visible;\nphysical AU/km values remain in the HUD.',618,395,310,16)
    d.text('Surface-only spin · parent compensation · quaternion rigs · logarithmic depth',32,462,570,15,True)
    d.end()

    d.start('Every motion has an owner and a clock','07 MOTION · Dated flight, moon revolution, axial spin, manual six-DOF and comet orientation',
        "")
    d.pic(TECH/'moon_motion.png',32,96,439,122);d.pic(TECH/'spin.png',489,96,439,122)
    d.text('MOON REVOLUTION / tilted parent frame',32,230,439,17,True)
    d.text('AXIAL SPIN / surface transform only',489,230,439,17,True)
    d.eq(r'\mathbf p(s)=h_{00}\mathbf p_0+h_{10}h\mathbf v_0+h_{01}\mathbf p_1+h_{11}h\mathbf v_1',32,270,896,22)
    d.eq(r'r=\frac{a(1-e^2)}{1+e\cos\nu},\quad\dot\nu=2\pi(0.4)/P_{days}',32,322,437,19)
    d.eq(r'\dot\psi=2\pi/P_{hours}',489,322,437,21)
    d.eq(r'q^{\prime}=\operatorname{normalize}(q q_y q_x q_z)',32,376,437,19)
    d.eq(r'\mathbf v^{\prime}=\mathbf v+\mathbf a\Delta t,\quad\mathbf p^{\prime}=\mathbf p+\mathbf v\Delta t',489,376,437,20)
    d.text('Historical data: one date. Visual moons/spin: independent rates.\nManual flight + cameras: real time. Comet tail: anti-Sun direction.',32,435,896,18,True)
    d.end()

    d.start('Three light types—and a moving headlamp','08 LIGHTING · Source geometry and attenuation determine incident illumination',
        "")
    d.pic(TECH/'lighting.png',32,92,896,193)
    for i,(title,formula) in enumerate([('POINT / Sun',r'\mathbf l=(\mathbf p_L-\mathbf p)/d'),('SPOT / camera',r'C=\mathrm{clamp}\left(\frac{\theta-\cos18^\circ}{\cos12^\circ-\cos18^\circ},0,1\right)'),('DIRECTIONAL / fill',r'\mathbf l=-\operatorname{normalize}(\mathbf d_L)')]):
        x=32+i*306;d.text(title,x,295,286,18,True,color=TEAL);d.eq(formula,x,330,286,18)
    d.eq(r'A(d)=\frac{1}{k_c+k_ld+k_qd^2},\quad D=\max(\mathbf n\cdot\mathbf l,0)',32,389,535,21)
    d.pic(TECH/'moving_spot.png',606,391,322,82)
    d.text('Camera motion carries the spotlight position and direction; the cone is 12°–18°.',32,472,896,16)
    d.end()

    d.start('Five shading techniques, matched views','09 SHADING · Change the normal/evaluation stage while keeping geometry, texture and lights fixed',
        "")
    d.pic(TECH/'shading.png',32,94,896,137)
    d.pic(TECH/'shading_voyager.png',32,230,896,137)
    modes=[('FLAT','Face normal'),('GOURAUD','Vertex light'),('PHONG','Pixel reflect'),('BLINN','Pixel half-vector'),('TOON','Bands + rim')]
    for i,(a,b) in enumerate(modes):
        x=32+i*181;d.text(a,x,372,172,15,True,color=TEAL);d.text(b,x,396,172,14)
    d.eq(r'S_P=k_s\max(\mathbf v\cdot\mathbf r,0)^p',32,437,436,20)
    d.eq(r'S_B=k_s\max(\mathbf n\cdot\mathbf h,0)^{2p}',489,437,436,20)
    d.text('Toon: 1 / 0.62 / 0.30 / 0.06 bands; rim at n·v < 0.22. Glow/unlit are separate materials.',32,478,896,13)
    d.end()

    d.start('Visibility: eclipses, ring shadows and self-shadow','10 SHADOWS · Analytic Sun rays for bodies/rings; a depth map for the spacecraft',
        "")
    d.pic(TECH/'shadows.png',32,94,896,203)
    d.eq(r'V_{soft}=1-\frac{A_{overlap}}{\pi\alpha_{Sun}^{2}},\quad V_{ring}=1-\mathrm{opacity}',32,318,896,23)
    d.rect(32,385,437,99);d.rect(489,385,439,99)
    d.text('BODIES + RINGS',46,395,407,19,True,color=TEAL)
    d.text('Analytic shadow rays: sphere / annulus.\nFinite Sun discs create a stable penumbra.',46,429,407,16)
    d.text('VOYAGER SELF-SHADOW',503,395,410,19,True,color=TEAL)
    d.text('2048² depth map · 3×3 PCF comparisons.\nBias avoids the surface shadowing itself.',503,429,410,16)
    d.end()

    d.start('Ray tracing: nearest hit, light, reflection','11 RAY TRACING · Analytic spheres/rings and a triangle BVH for the authored spacecraft',
        "")
    d.pic(TECH/'tracing.png',32,90,896,152)
    d.pic(TECH/'ray_paths.png',32,250,438,179)
    d.pic(TECH/'bvh_layout.png',489,250,439,179)
    d.eq(r'\mathbf p(t)=\mathbf o+t\mathbf d,\quad t=-b\pm\sqrt{b^2-f}',32,437,438,18)
    d.eq(r'\mathbf d_r=\mathbf d-2(\mathbf d\cdot\mathbf n)\mathbf n',489,437,438,20)
    d.text('F9 hybrid Whitted view · 5,828 BVH triangles · up to 4 ring layers · F10 one mirror bounce',32,478,896,14)
    d.end()

    d.start('Textures: colour, atlas, relief and specular mask','12 TEXTURES · The mesh supplies shape; image channels supply surface detail',
        "")
    d.pic(TECH/'texture_channels.png',32,93,896,160)
    d.pic(TECH/'surface_maps.png',32,269,436,159)
    d.pic(ROOT/'assets/textures/spacecraft/voyager_nasa_atlas.png',491,269,209,159)
    d.text('NASA atlas\n11 textured finishes\n+ copper material',719,294,208,17,True)
    d.eq(r'\mathbf n_t=\operatorname{normalize}(-sS_x(H),-sS_y(H),1)',32,443,595,21)
    d.eq(r'uv^{\prime}=uv\cdot zw+xy',651,443,276,21)
    d.text('F8 maps off/on · Sobel relief is inferred from colour · ocean mask controls highlights',32,479,896,14)
    d.end()

    d.start('Thank you','13 DISCUSSION · Questions about the project, its mathematics and its implementation',
        "")
    d.text('From indexed triangles\nto an explorable mission.',32,113,525,32,True)
    d.text('Geometry · transforms · motion\nlighting · shading · ray tracing · textures',32,208,525,22,color=TEAL)
    d.pic(FIG/'jupiter_encounter.jpg',585,104,343,255)
    d.rect(32,343,525,126)
    d.text('Verified: builds, scene checks and runtime captures.\nMeasured raster views: 0.64–1.15 ms / RX 590.\nExplainable limits: compressed scale, visual moons,\nhybrid tracing and one reflection bounce.',47,358,495,18,leading=25)
    d.text('Sarwad Hasan Siddiqui\nRoll: 2107006 | CSE 4102 | KUET',586,380,339,19,True)
    d.end();d.save()
    print('[SLIDES] 13 slides: 2 introductions + 10 technical + thank you')

if __name__=='__main__':main()
