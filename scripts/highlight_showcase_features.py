"""Clarify feature captions in edited presentations without rebuilding slides."""
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import shutil
import zipfile

import pymupdf
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
SUBTITLES = {
    2: '02 INTRODUCTION · Important features, screenshots and narrated project demo',
    3: '03 GEOMETRY · One indexed sphere reused for 26 textured bodies; vertices, normals and UVs define the surface',
    4: '04 SPACECRAFT · Authored boxes, rods and a curved dish create 76 assemblies and 18 inspection targets',
    5: '05 ENVIRONMENT · Four ring systems, 8,500 instanced rocks, 7,380 stars and a comet complete the scene',
    6: '06 INTERACTION · Four camera modes and floating origin let you inspect planets and metre-scale hardware',
    7: '07 MOTION · NASA/JPL dated flybys, axial spin, moon revolution and user-controlled spacecraft flight',
    8: '08 LIGHTING · Sun, moving camera spotlight and directional fill; compare their responses in the same view',
    9: '09 SHADING · Compare Flat, Gouraud, Phong, Blinn–Phong and Toon on the same planets and spacecraft',
    10: '10 SHADOWS · Analytic rays show eclipses and ring shadows; a depth map handles spacecraft self-shadow',
    11: '11 BONUS · Hybrid tracing uses analytic planets/rings, a spacecraft triangle BVH and one optional reflection',
    12: '12 SURFACES · NASA atlas and derived normal/specular maps add detail; F8 toggles maps for comparison',
}
INTRO_TITLE = 'Important features: explore, inspect, compare'
INTRO_CAPTIONS = [
    ('MISSION', 'DATED MISSION', 'Dated launch, four flybys and\nheliopause.',
     'NASA/JPL data on one date.\nReplay four giant-planet flybys.'),
    ('INSPECTION', 'PROCEDURAL VOYAGER', 'Body focus, 18 hardware targets, free\ncamera.',
     '76 authored parts; 18 targets.\nInspect curved dish + NASA atlas.'),
    ('COMPARISON', 'RENDERING LAB', 'Three lights, five shading modes, ray\ntracing.',
     'Three lights; five shading modes.\nBVH tracing + one mirror bounce.'),
]
FEATURE_NOTES = '''IMPORTANT FEATURES AND VIDEO NARRATION
Use the three screenshots above the captions as evidence. Begin with the dated mission: NASA/JPL offline positions and velocities place planets and Voyager on one shared Julian Date; interpolation gives smooth flybys. In the video, point to the running date and the four giant-planet encounters. This connects historical mission data to animation and transforms (slides 2 and 7).

Next describe the spacecraft. The mesh is authored from reusable primitives rather than an imported spacecraft mesh: boxes, cylinders, rods and a sampled paraboloid. There are 76 assemblies and 18 inspectable hardware targets. Point to the dish curvature, truss structure, component labels and NASA texture atlas. The model and inspection system connect geometry, normals, hierarchical transforms and interaction (slides 2, 4, 6 and 12).

Show user control: four camera modes, mouse selection, zoom and real-time manual six-degree flight. Floating origin subtracts camera position in double precision so hardware close-ups remain stable across large world coordinates. In the video show a selected body, then Inspect, then thrust/yaw/braking; explain that input changes state and the update computes the new pose (slides 6–7).

Show the rendering comparison in a frozen view. Toggle the moving camera spotlight and directional fill, then compare Gouraud against Phong and the other three shading modes. Identify the visible response and where illumination is computed; explain ambient, Lambert diffuse and the specular exponent rather than only naming techniques (slides 8–9).

The bonus rendering feature is hybrid Whitted ray tracing. F9 traces analytic planet/ring hits and the authored spacecraft through a triangle BVH. F10 enables one optional mirror reflection. Show a raster/traced comparison and describe primary, shadow and reflected rays, then show body/ring shadows and spacecraft self-shadow (slides 10–11).

Texture detail is another important feature: the NASA atlas supplies spacecraft finishes; colour-derived relief and an ocean mask affect normal/specular response. F8 compares maps off/on without changing geometry. Explain that normal maps change shading normals, not the silhouette (slide 12).

Finally explain the environment and efficiency: four ring systems, 8,500 deterministic rock instances in three fields, 7,380 star points and a dynamic comet. Instancing, shared sphere meshes, LOD and culling reduce drawing work. Rock fields are static; the comet and mission/body animation supply motion (slides 3 and 5).

Narrate the video using feature → calculation → visible effect. Use the two-minute plan in docs/guide/13-teacher-showcase.md. The project-specific strengths are the integration, authored geometry and inspectable comparisons; the graphics algorithms themselves are established techniques.
'''

def replace_text(shape, text):
    frame = shape.text_frame
    templates = list(frame.paragraphs)
    first = templates[0]
    run_style = deepcopy(first.runs[0]._r.rPr) if first.runs and first.runs[0]._r.rPr is not None else None
    for child in list(frame._txBody):
        if child.tag.endswith('}p'):
            frame._txBody.remove(child)
    for i, line in enumerate(text.split('\n')):
        template = templates[min(i, len(templates)-1)]
        paragraph_xml = deepcopy(template._p)
        for child in list(paragraph_xml):
            if not child.tag.endswith('}pPr') and not child.tag.endswith('}endParaRPr'):
                paragraph_xml.remove(child)
        frame._txBody.append(paragraph_xml)
        paragraph = frame.paragraphs[-1]
        run = paragraph.add_run()
        run.text = line
        if run_style is not None:
            run._r.insert(0, deepcopy(run_style))

def update_deck(path):
    original = path.read_bytes()
    ppt = Presentation(BytesIO(original))
    if len(ppt.slides) != 13:
        raise ValueError('Expected the current thirteen-slide edited deck')
    changes = []
    targets = set()
    for number, subtitle in SUBTITLES.items():
        slide = ppt.slides[number-1]
        candidates = [s for s in slide.shapes if s.has_text_frame and s.top < 2*914400
                      and s.text.startswith(f'{number:02} ')]
        if len(candidates) != 1:
            raise ValueError(f'Cannot identify subtitle on slide {number}')
        shape = candidates[0]
        if shape.text == subtitle:
            continue
        changes.append({'slide':number, 'old':shape.text, 'new':subtitle,
                        'box':[float(v)/12700 for v in (shape.left,shape.top,shape.width,shape.height)]})
        replace_text(shape, subtitle)
        targets.add(str(slide.part.partname).lstrip('/'))
    slide = ppt.slides[1]
    replacements = {'Explore, control, compare':INTRO_TITLE}
    for old_heading, heading, old_body, body in INTRO_CAPTIONS:
        replacements[old_heading] = heading
        replacements[old_body] = body
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text in replacements:
            old = shape.text
            new = replacements[old]
            changes.append({'slide':2,'old':old,'new':new,
                            'box':[float(v)/12700 for v in (shape.left,shape.top,shape.width,shape.height)]})
            replace_text(shape,new)
    notes = slide.notes_slide.notes_text_frame
    if FEATURE_NOTES.splitlines()[0] not in notes.text:
        notes.text = notes.text.rstrip()+'\n\n'+FEATURE_NOTES
    elif not changes:
        return []
    targets.add(str(slide.notes_slide.part.partname).lstrip('/'))
    generated = BytesIO()
    ppt.save(generated)
    output = BytesIO()
    with zipfile.ZipFile(BytesIO(original)) as source, zipfile.ZipFile(generated) as modified, zipfile.ZipFile(output,'w') as dest:
        for item in source.infolist():
            dest.writestr(item,modified.read(item.filename) if item.filename in targets else source.read(item.filename))
    with zipfile.ZipFile(BytesIO(original)) as source, zipfile.ZipFile(output) as modified:
        assert source.namelist()==modified.namelist()
        assert all(source.read(name)==modified.read(name) for name in source.namelist() if name not in targets)
    path.write_bytes(output.getvalue())
    return changes

def update_pdf(changes):
    if not changes:
        return
    path = ROOT/'docs/report/Voyager_2_Presentation.pdf'
    doc = pymupdf.open(path)
    fonts = {'bold':pymupdf.Font(fontfile='C:/Windows/Fonts/arialbd.ttf'),
             'regular':pymupdf.Font(fontfile='C:/Windows/Fonts/arial.ttf')}
    for number in sorted({item['slide'] for item in changes}):
        page = doc[number-1]
        edits = []
        for item in changes:
            if item['slide'] != number:
                continue
            old = item['old'].replace('\n',' ')
            x,y,w,h = item['box']
            authored_box = pymupdf.Rect(x,y,x+w,y+h+3)
            hits = [hit for hit in page.search_for(old) if hit.intersects(authored_box)]
            if not hits:
                raise ValueError(f"Cannot locate PDF text on slide {number}: {old}")
            spans = [span for block in page.get_text('dict')['blocks'] for line in block.get('lines',[])
                     for span in line['spans'] if pymupdf.Rect(span['bbox']).intersects(hits[0])]
            span = spans[0]
            for hit in hits:
                page.add_redact_annot(hit, fill=(1,1,1))
            edits.append((item,pymupdf.Rect(x,y,x+w,y+h+3),span))
        page.apply_redactions(images=0,graphics=0)
        for item,rect,span in edits:
            font = fonts['bold' if 'Bold' in span['font'] else 'regular']
            name = 'featureBold' if 'Bold' in span['font'] else 'featureRegular'
            page.insert_font(fontname=name,fontbuffer=font.buffer)
            colour = ((span['color']>>16 & 255)/255,(span['color']>>8 & 255)/255,(span['color'] & 255)/255)
            if page.insert_textbox(rect,item['new'],fontname=name,fontsize=span['size'],lineheight=1.16,color=colour)<0:
                raise ValueError(f"Feature description overflows on slide {number}: {item['new']}")
    temporary = path.with_name('Voyager_2_Presentation.features.pdf')
    doc.save(temporary,garbage=3,deflate=True)
    doc.close()
    temporary.replace(path)

def main():
    backup = ROOT/'captures/assessment/showcase/features-before'
    backup.mkdir(parents=True,exist_ok=True)
    names = ['Voyager_2_Presentation.pptx','Final-Presentation.pptx']
    for name in names:
        src = ROOT/'docs/report'/name
        if not (backup/name).exists():
            shutil.copy2(src,backup/name)
    pdf = ROOT/'docs/report/Voyager_2_Presentation.pdf'
    if not (backup/pdf.name).exists():
        shutil.copy2(pdf,backup/pdf.name)
    changes = [update_deck(ROOT/'docs/report'/name) for name in names]
    update_pdf(changes[0])
    (ROOT/'captures/assessment/showcase/feature-text-changes.json').write_text(json.dumps(changes,indent=2)+'\n',encoding='utf-8')
    print('Feature captions and narration added; existing slide layouts and media preserved.')

if __name__=='__main__':
    main()
